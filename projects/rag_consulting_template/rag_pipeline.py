"""
rag_pipeline.py — a generic, document-agnostic Retrieval-Augmented Generation
pipeline.

Point it at any folder of documents (.txt, .md, .csv, .pdf, .docx) and it
will chunk, embed, index, retrieve, and answer questions grounded in that
folder's content. No assumptions about domain, language, or document
structure are baked in, unlike ontologyrag_q_kaggle (which is specific to
the Tafsir/Quran dataset) — this is meant to be handed a new client's
document set and reused as-is.

Usage as a library:

    from rag_pipeline import RAGPipeline

    rag = RAGPipeline()
    rag.build_index("path/to/documents")
    answer, sources = rag.answer("What is the refund policy?")

Usage as a CLI:

    python rag_pipeline.py --docs ./my_docs --query "What is the refund policy?"
    python rag_pipeline.py --docs ./my_docs --interactive
"""

from __future__ import annotations

import argparse
import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional


# ---------------------------------------------------------------------------
# 1. Document loading
# ---------------------------------------------------------------------------

SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".pdf", ".docx"}


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_csv(path: Path) -> str:
    import csv

    lines = []
    with open(path, encoding="utf-8", errors="ignore") as f:
        for row in csv.reader(f):
            lines.append(" | ".join(row))
    return "\n".join(lines)


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ImportError("Reading .pdf files requires: pip install pypdf") from e
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as e:
        raise ImportError("Reading .docx files requires: pip install python-docx") from e
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


_READERS = {
    ".txt": _read_txt,
    ".md": _read_txt,
    ".csv": _read_csv,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
}


@dataclass
class Document:
    source: str  # file path, relative to the docs folder
    text: str


def load_documents(folder: str | Path) -> list[Document]:
    """Recursively load every supported file under `folder` into Documents."""
    folder = Path(folder)
    if not folder.is_dir():
        raise NotADirectoryError(f"{folder} is not a directory")

    docs = []
    for path in sorted(folder.rglob("*")):
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            text = _READERS[path.suffix.lower()](path)
        except Exception as e:
            print(f"Skipping {path} ({e})")
            continue
        text = text.strip()
        if text:
            docs.append(Document(source=str(path.relative_to(folder)), text=text))
    return docs


# ---------------------------------------------------------------------------
# 2. Chunking
# ---------------------------------------------------------------------------

@dataclass
class Chunk:
    chunk_id: int
    source: str
    text: str


def chunk_text(text: str, chunk_size: int = 200, overlap: int = 40) -> list[str]:
    """Word-count sliding-window chunking. Simple and language-agnostic --
    works for Arabic, English, or any whitespace-delimited language.
    """
    words = text.split()
    if not words:
        return []
    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size")

    chunks = []
    start = 0
    step = chunk_size - overlap
    while start < len(words):
        piece = words[start : start + chunk_size]
        chunks.append(" ".join(piece))
        if start + chunk_size >= len(words):
            break
        start += step
    return chunks


def build_chunks(documents: list[Document], chunk_size: int = 200, overlap: int = 40) -> list[Chunk]:
    chunks = []
    next_id = 0
    for doc in documents:
        for piece in chunk_text(doc.text, chunk_size=chunk_size, overlap=overlap):
            chunks.append(Chunk(chunk_id=next_id, source=doc.source, text=piece))
            next_id += 1
    return chunks


# ---------------------------------------------------------------------------
# 3. The pipeline: embed, index, retrieve, generate
# ---------------------------------------------------------------------------

DEFAULT_EMBEDDING_MODEL = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
DEFAULT_LOCAL_LLM = "Qwen/Qwen2.5-1.5B-Instruct"

PROMPT_TEMPLATE = (
    "Answer the question using ONLY the information in the context below. "
    "If the context doesn't contain the answer, say you don't know.\n\n"
    "Context:\n{context}\n\nQuestion: {question}\nAnswer:"
)


class RAGPipeline:
    def __init__(
        self,
        embedding_model: str = DEFAULT_EMBEDDING_MODEL,
        openai_api_key: Optional[str] = None,
        openai_model: str = "gpt-4o-mini",
        local_llm: str = DEFAULT_LOCAL_LLM,
    ):
        self.embedding_model_name = embedding_model
        self.openai_api_key = openai_api_key
        self.openai_model = openai_model
        self.local_llm_name = local_llm

        self._embedder = None
        self._index = None
        self._chunks: list[Chunk] = []
        self._gen_client = None  # openai client or (tokenizer, model) tuple

    # -- lazy loaders, so importing this module never triggers a download --
    def _get_embedder(self):
        if self._embedder is None:
            from sentence_transformers import SentenceTransformer

            self._embedder = SentenceTransformer(self.embedding_model_name)
        return self._embedder

    def _get_generator(self):
        if self._gen_client is not None:
            return self._gen_client

        if self.openai_api_key:
            from openai import OpenAI

            self._gen_client = ("openai", OpenAI(api_key=self.openai_api_key))
        else:
            import torch
            from transformers import AutoModelForCausalLM, AutoTokenizer

            tokenizer = AutoTokenizer.from_pretrained(self.local_llm_name)
            model = AutoModelForCausalLM.from_pretrained(
                self.local_llm_name,
                torch_dtype=torch.float16 if torch.cuda.is_available() else torch.float32,
                device_map="auto",
            )
            self._gen_client = ("local", (tokenizer, model))
        return self._gen_client

    # -- indexing --
    def build_index(self, folder: str | Path, chunk_size: int = 200, overlap: int = 40) -> None:
        documents = load_documents(folder)
        if not documents:
            raise ValueError(f"No supported documents found under {folder}")
        self._chunks = build_chunks(documents, chunk_size=chunk_size, overlap=overlap)
        if not self._chunks:
            raise ValueError("Documents were loaded but produced zero chunks")

        import numpy as np
        import faiss

        embedder = self._get_embedder()
        texts = [c.text for c in self._chunks]
        embeddings = embedder.encode(texts, batch_size=64, show_progress_bar=True, normalize_embeddings=True)
        self._index = faiss.IndexFlatIP(embeddings.shape[1])
        self._index.add(np.asarray(embeddings, dtype="float32"))

    def save_index(self, path: str | Path) -> None:
        """Persist the FAISS index + chunk metadata so you don't have to
        re-embed the corpus on every run (useful once a client's document
        set is large)."""
        if self._index is None:
            raise RuntimeError("Nothing to save -- call build_index() first")
        import faiss

        path = Path(path)
        path.mkdir(parents=True, exist_ok=True)
        faiss.write_index(self._index, str(path / "index.faiss"))
        with open(path / "chunks.json", "w", encoding="utf-8") as f:
            json.dump([c.__dict__ for c in self._chunks], f, ensure_ascii=False)

    def load_index(self, path: str | Path) -> None:
        import faiss

        path = Path(path)
        self._index = faiss.read_index(str(path / "index.faiss"))
        with open(path / "chunks.json", encoding="utf-8") as f:
            self._chunks = [Chunk(**c) for c in json.load(f)]

    # -- retrieval --
    def retrieve(self, query: str, top_k: int = 3) -> list[Chunk]:
        if self._index is None:
            raise RuntimeError("Call build_index() or load_index() first")
        import numpy as np

        embedder = self._get_embedder()
        q_emb = embedder.encode([query], normalize_embeddings=True)
        _, idxs = self._index.search(np.asarray(q_emb, dtype="float32"), top_k)
        return [self._chunks[i] for i in idxs[0] if i != -1]

    # -- generation --
    def _generate(self, question: str, context: str) -> str:
        prompt = PROMPT_TEMPLATE.format(context=context, question=question)
        kind, client = self._get_generator()

        if kind == "openai":
            resp = client.chat.completions.create(
                model=self.openai_model,
                messages=[{"role": "user", "content": prompt}],
                temperature=0.2,
            )
            return resp.choices[0].message.content.strip()

        tokenizer, model = client
        messages = [{"role": "user", "content": prompt}]
        text = tokenizer.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
        inputs = tokenizer(text, return_tensors="pt").to(model.device)
        output = model.generate(**inputs, max_new_tokens=250, do_sample=False)
        answer = tokenizer.decode(output[0][inputs["input_ids"].shape[1] :], skip_special_tokens=True)
        return answer.strip()

    def answer(self, question: str, top_k: int = 3) -> tuple[str, list[Chunk]]:
        """The full RAG call: retrieve relevant chunks, then generate a
        grounded answer. Returns (answer_text, retrieved_chunks) so callers
        can show their sources."""
        retrieved = self.retrieve(question, top_k=top_k)
        context = "\n---\n".join(f"[{c.source}]\n{c.text}" for c in retrieved)
        answer_text = self._generate(question, context)
        return answer_text, retrieved


# ---------------------------------------------------------------------------
# 4. CLI
# ---------------------------------------------------------------------------

def _main():
    parser = argparse.ArgumentParser(description="Generic RAG pipeline over a folder of documents")
    parser.add_argument("--docs", required=True, help="Folder of documents to index")
    parser.add_argument("--query", help="A single question to ask")
    parser.add_argument("--interactive", action="store_true", help="Ask multiple questions in a loop")
    parser.add_argument("--top-k", type=int, default=3)
    parser.add_argument("--chunk-size", type=int, default=200)
    parser.add_argument("--overlap", type=int, default=40)
    parser.add_argument("--openai-key", default=os.environ.get("OPENAI_API_KEY"))
    parser.add_argument("--save-index", help="Directory to persist the index to after building")
    parser.add_argument("--load-index", help="Directory to load a previously saved index from")
    args = parser.parse_args()

    rag = RAGPipeline(openai_api_key=args.openai_key)

    if args.load_index:
        rag.load_index(args.load_index)
    else:
        rag.build_index(args.docs, chunk_size=args.chunk_size, overlap=args.overlap)
        if args.save_index:
            rag.save_index(args.save_index)

    def ask(question: str):
        answer_text, sources = rag.answer(question, top_k=args.top_k)
        print("\nQ:", question)
        print("A:", answer_text)
        print("Sources:", sorted({c.source for c in sources}))

    if args.query:
        ask(args.query)
    if args.interactive or not args.query:
        print("\nEnter questions (empty line to quit):")
        while True:
            q = input("> ").strip()
            if not q:
                break
            ask(q)


if __name__ == "__main__":
    _main()
